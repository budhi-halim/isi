"""Exercise unchanged workbook generators with synthetic input and inspect their output."""

from datetime import datetime
from io import BytesIO
from pathlib import Path
import json
import runpy
import sys
from types import ModuleType, SimpleNamespace

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = Path(__file__).resolve().parent / "evidence" / "checkpoint-6-7"


class BrowserArray(list):
    """Model the existing JavaScript array boundary consumed by PyScript."""

    def to_py(self) -> list:
        """Return the input as a normal Python list."""
        return list(self)


def fixture(year: int = 2026) -> bytes:
    """Create a small workbook using the established input columns and header row."""
    workbook = Workbook()
    sheet = workbook.active
    headers = ["Area", "Date", "Invoice No", "Customer Name", "Product Type", "Product Name", "Quantity", "Unit Price"]
    for column, header in enumerate(headers, 2):
        sheet.cell(3, column, header)
    for row, (area, month, quantity, price) in enumerate([
        ("Jakarta", 1, 10, 15000), ("Bandung", 2, 25, 20000), ("Jakarta", 3, 15, 18000)
    ], 4):
        values = [area, datetime(year, month, 15), f"TEST-{row}", "Fixture Customer", "Spices", "Fixture Product", quantity, price]
        for column, value in enumerate(values, 2):
            sheet.cell(row, column, value)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def load_generator(repo: str) -> dict:
    """Load an unchanged application module with the browser bridge stubbed."""
    bridge = ModuleType("pyscript")
    bridge.ffi = SimpleNamespace(create_proxy=lambda value: value)
    bridge.window = SimpleNamespace()
    message = SimpleNamespace(innerHTML="")
    bridge.document = SimpleNamespace(getElementById=lambda name: message)
    sys.modules["pyscript"] = bridge
    return runpy.run_path(str(ROOT / repo / "py" / "main.py"))


def inspect(buffer: list[int], repo: str) -> dict:
    """Check data, formulas, cell formats, tables, and target chart preservation."""
    workbook = load_workbook(BytesIO(bytes(buffer)))
    sales = workbook["Sales Data"]
    assert sales.max_row == 4
    assert [sales.cell(row, 8).value for row in range(2, 5)] == [10, 25, 15]
    assert [sales.cell(row, 9).value for row in range(2, 5)] == [15000, 20000, 18000]
    assert sales["J2"].value == "=PRODUCT(H2,I2)"
    assert sales["B2"].number_format == "dd/mm/yyyy"
    assert sales["H2"].number_format == "#,##0"
    assert "Sales_Data" in sales.tables
    assert workbook["Exchange Rate"].max_row == 4
    if repo == "sales-target":
        assert workbook["Cumulative Percentage"].sheet_state == "veryHidden"
        assert len(workbook["Sales Target"]._charts) == 2
    (OUTPUT / f"{repo}-fixture-output.xlsx").write_bytes(bytes(buffer))
    return {"sheets": workbook.sheetnames, "sales_rows": sales.max_row - 1, "tables": sum(len(sheet.tables) for sheet in workbook)}


def compare_browser(repo: str) -> int:
    """Compare browser output to direct Python output without volatile ZIP metadata."""
    expected = load_workbook(OUTPUT / f"{repo}-fixture-output.xlsx")
    actual = load_workbook(OUTPUT / f"{repo}-browser-output.xlsx")
    assert expected.sheetnames == actual.sheetnames
    count = 0
    for left, right in zip(expected, actual):
        assert (left.sheet_state, left.max_row, left.max_column, str(left.merged_cells), len(left._charts), list(left.tables)) == (
            right.sheet_state, right.max_row, right.max_column, str(right.merged_cells), len(right._charts), list(right.tables)
        )
        for left_row, right_row in zip(left, right):
            for left_cell, right_cell in zip(left_row, right_row):
                assert (left_cell.value, left_cell.data_type, left_cell.number_format, left_cell.style_id) == (
                    right_cell.value, right_cell.data_type, right_cell.number_format, right_cell.style_id
                ), (repo, left.title, left_cell.coordinate)
                count += 1
    return count


def main() -> None:
    """Run each public entrypoint with normal, empty, and multiple-year fixtures."""
    OUTPUT.mkdir(parents=True, exist_ok=True)
    data = fixture()
    (OUTPUT / "input-fixture.xlsx").write_bytes(data)
    (OUTPUT / "input-fixture-2025.xlsx").write_bytes(fixture(2025))
    files = BrowserArray([{"name": "fixture.xlsx", "data": list(data)}])
    report = load_generator("sales-report")
    result = report["process_files"](files)
    assert result["type"] == "success", result.get("message")
    summary = {"sales-report": inspect(result["buffer"], "sales-report")}
    assert "buffer" not in report["process_files"](BrowserArray())
    target = load_generator("sales-target")
    prepared = target["prepare_files"](files)
    assert prepared["multi_year"] is False
    assert len(prepared["areas"]) == 2
    result = target["finalize_files"](files, BrowserArray([list(reversed(prepared["areas"]))]), True)
    assert result["type"] == "success", result.get("message")
    summary["sales-target"] = inspect(result["buffer"], "sales-target")
    multiple = BrowserArray([*files, {"name": "prior.xlsx", "data": list(fixture(2025))}])
    assert target["prepare_files"](multiple)["multi_year"] is True
    assert target["finalize_files"](multiple, BrowserArray([]), False)["type"] == "error"
    assert target["prepare_files"](BrowserArray())["type"] == "error"
    if "--compare-browser" in sys.argv:
        for repo in ["sales-report", "sales-target"]:
            summary[repo]["browser_cells_matched"] = compare_browser(repo)
    (OUTPUT / "workbook-results.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary))


if __name__ == "__main__":
    main()
